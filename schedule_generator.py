"""
상담센터 월간 근무표 생성기

매달 상단 설정값(연도/월, 명단, 예외일 등)만 바꿔서 다시 실행하면
새로운 근무표(.xlsx)가 생성된다.
"""

from __future__ import annotations

import calendar
import json
import os
from dataclasses import dataclass, field
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# =============================================================================
# 설정값 (매달 이 블록만 수정해서 재사용)
# =============================================================================

# 작업 스케줄러 등으로 자동 실행할 때는 True로 두면 실행 시점 기준 "다음 달"
# 근무표를 자동으로 생성한다. 특정 달을 지정해서 수동으로 돌리고 싶으면
# False로 바꾸고 아래 YEAR/MONTH를 원하는 값으로 지정한다.
AUTO_NEXT_MONTH = True

YEAR = 2026
MONTH = 11

if AUTO_NEXT_MONTH:
    _today = date.today()
    YEAR, MONTH = (_today.year + 1, 1) if _today.month == 12 else (_today.year, _today.month + 1)

# 전체 상담사 명단 (1번/2번 좌석 배정 대상이 되는 전원)
ALL_COUNSELORS = [
    "김병성", "박민선", "김민성", "박미나",
    "이민정", "김윤정", "이영숙",
    "추의성", "김은숙",
]

# 1번 좌석 제외 대상 (시니어일자리지원센터 컨설턴트) - 2번 좌석 근무는 가능
SEAT1_EXCLUDED = ["추의성", "김은숙"]

# 상담사별 근무 가능 요일 (0=월 ~ 4=금). 지정하지 않은 사람은 평일 전체 근무 가능으로 간주.
WEEKLY_AVAILABILITY: dict[str, set[int]] = {
    # 예) "박미나": {0, 1, 2},  # 월,화,수만 근무
}

# 특정 날짜에만 개별적으로 근무 불가능한 경우 (휴가 등)
AD_HOC_LEAVE: dict[str, set[date]] = {
    # 예) "이민정": {date(2026, 9, 15)},
}

# 공휴일 / 특별일정 (해당 기간은 근무 배정에서 제외하고, 사유 문구를 표시)
# 자동 실행(AUTO_NEXT_MONTH=True) 전에는 이번 달 것으로 반드시 업데이트할 것 —
# 공휴일/행사일은 자동으로 알 수 없어 유일하게 매달 손으로 채워야 하는 항목이다.
SPECIAL_SCHEDULES: list[dict] = [
    {"start": date(2026, 10, 9), "end": date(2026, 10, 9), "label": "한글날"},  # 예시 (10월)
]

# 전월까지의 1번 좌석 누적 배정 횟수 / 마지막 배정일을 여기 직접 채워도 되지만,
# 보통은 비워 두면 된다 — 매달 실행이 끝날 때마다 STATE_FILE에 자동 저장되고,
# 다음 실행 때 자동으로 이어받아 형평성을 유지한다. 여기 값을 채우면 그 값이
# STATE_FILE의 값보다 우선한다(예: 담당자가 바뀌어 수동으로 다시 맞춰야 할 때).
PREV_SEAT1_COUNTS: dict[str, int] = {
    # 예) "김병성": 3,
}
PREV_SEAT1_LAST_DATE: dict[str, date] = {
    # 예) "김병성": date(2026, 9, 22),
}

# 상담사별 1번 좌석 누적 배정 현황을 저장해 두는 파일 (자동 실행 시 형평성 이어받기용)
STATE_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "schedule_state.json")

OUTPUT_FILENAME = f"{YEAR}년_{MONTH}월_상담센터_근무표.xlsx"

WEEKDAY_LABELS = ["월", "화", "수", "목", "금"]


# =============================================================================
# 데이터 구조
# =============================================================================

@dataclass
class DaySchedule:
    seat1: str | None = None
    seat2: str | None = None
    note: str | None = None  # 예외일 사유 (공휴일/행사 등)


@dataclass
class ScheduleResult:
    year: int
    month: int
    days: dict[date, DaySchedule] = field(default_factory=dict)
    seat1_counts_this_month: dict[str, int] = field(default_factory=dict)
    # 다음 달 실행 시 형평성을 이어가기 위한 누적치(전월 값 + 이번 달 결과)
    seat1_counts_cumulative: dict[str, int] = field(default_factory=dict)
    seat1_last_date: dict[str, date] = field(default_factory=dict)


# =============================================================================
# 누적 배정 현황 저장/불러오기 (자동 실행 시 형평성 이어받기)
# =============================================================================

def load_state(state_file: str) -> tuple[dict[str, int], dict[str, date]]:
    if not os.path.exists(state_file):
        return {}, {}
    with open(state_file, encoding="utf-8") as f:
        data = json.load(f)
    counts = data.get("seat1_counts", {})
    last_date = {c: date.fromisoformat(d) for c, d in data.get("seat1_last_date", {}).items()}
    return counts, last_date


def save_state(state_file: str, counts: dict[str, int], last_date: dict[str, date]) -> None:
    data = {
        "seat1_counts": counts,
        "seat1_last_date": {c: d.isoformat() for c, d in last_date.items()},
    }
    with open(state_file, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


# =============================================================================
# 예외일 처리 로직
# =============================================================================

def build_exception_map(special_schedules: list[dict]) -> dict[date, str]:
    """공휴일/특별일정 목록을 '날짜 -> 표시문구' 매핑으로 펼친다."""
    exception_map: dict[date, str] = {}
    for item in special_schedules:
        d = item["start"]
        while d <= item["end"]:
            exception_map[d] = item["label"]
            d = date.fromordinal(d.toordinal() + 1)
    return exception_map


def get_weekdays_of_month(year: int, month: int) -> list[date]:
    """해당 월의 평일(월~금) 날짜 목록을 반환한다."""
    _, last_day = calendar.monthrange(year, month)
    all_days = [date(year, month, d) for d in range(1, last_day + 1)]
    return [d for d in all_days if d.weekday() < 5]


# =============================================================================
# 형평성(1번 좌석 균등 배정) 로직
# =============================================================================

def is_available(counselor: str, target_date: date) -> bool:
    """해당 상담사가 해당 날짜에 근무 가능한지 판단한다."""
    allowed_weekdays = WEEKLY_AVAILABILITY.get(counselor)
    if allowed_weekdays is not None and target_date.weekday() not in allowed_weekdays:
        return False
    if target_date in AD_HOC_LEAVE.get(counselor, set()):
        return False
    return True


def get_available_counselors(target_date: date, roster: list[str]) -> list[str]:
    return [c for c in roster if is_available(c, target_date)]


def pick_seat1(
    candidates: list[str],
    seat1_counts: dict[str, int],
    seat1_last_date: dict[str, date],
) -> str:
    """1번 좌석 배정 횟수가 가장 적은 사람 우선, 동률이면 가장 오래전에
    1번 좌석에 앉았던 사람(한 번도 없으면 최우선)을 선택한다."""
    return min(
        candidates,
        key=lambda c: (seat1_counts.get(c, 0), seat1_last_date.get(c, date.min)),
    )


def pick_seat2(candidates: list[str], total_counts: dict[str, int]) -> str:
    """전체 근무 횟수(1번+2번 합산)가 가장 적은 사람을 2번 좌석에 우선 배정해
    특정 인원에게 근무가 몰리지 않도록 한다."""
    return min(candidates, key=lambda c: total_counts.get(c, 0))


# =============================================================================
# 근무표 생성
# =============================================================================

def generate_schedule(
    year: int,
    month: int,
    all_counselors: list[str],
    seat1_excluded: list[str],
    special_schedules: list[dict],
    prev_seat1_counts: dict[str, int] | None = None,
    prev_seat1_last_date: dict[str, date] | None = None,
) -> ScheduleResult:
    exception_map = build_exception_map(special_schedules)
    weekdays = get_weekdays_of_month(year, month)

    # 형평성 판단에 쓰이는 누적치(전월 값 포함)
    seat1_counts_cumulative: dict[str, int] = dict(prev_seat1_counts or {})
    seat1_last_date: dict[str, date] = dict(prev_seat1_last_date or {})

    # 이번 달 결과 집계용
    seat1_counts_this_month: dict[str, int] = {c: 0 for c in all_counselors}
    total_counts_this_month: dict[str, int] = {c: 0 for c in all_counselors}

    result = ScheduleResult(year=year, month=month)

    for d in weekdays:
        if d in exception_map:
            result.days[d] = DaySchedule(note=exception_map[d])
            continue

        candidates = get_available_counselors(d, all_counselors)
        seat1_candidates = [c for c in candidates if c not in seat1_excluded]

        if not seat1_candidates:
            raise ValueError(f"{d}: 1번 좌석에 배정 가능한 상담사가 없습니다.")

        seat1_person = pick_seat1(seat1_candidates, seat1_counts_cumulative, seat1_last_date)

        seat2_candidates = [c for c in candidates if c != seat1_person]
        if not seat2_candidates:
            raise ValueError(f"{d}: 2번 좌석에 배정 가능한 상담사가 없습니다.")

        seat2_person = pick_seat2(seat2_candidates, total_counts_this_month)

        seat1_counts_cumulative[seat1_person] = seat1_counts_cumulative.get(seat1_person, 0) + 1
        seat1_last_date[seat1_person] = d
        seat1_counts_this_month[seat1_person] += 1
        total_counts_this_month[seat1_person] += 1
        total_counts_this_month[seat2_person] += 1

        result.days[d] = DaySchedule(seat1=seat1_person, seat2=seat2_person)

    result.seat1_counts_this_month = seat1_counts_this_month
    result.seat1_counts_cumulative = seat1_counts_cumulative
    result.seat1_last_date = seat1_last_date
    return result


# =============================================================================
# 자동 검증
# =============================================================================

def validate_schedule(
    result: ScheduleResult,
    all_counselors: list[str],
    seat1_excluded: list[str],
    special_schedules: list[dict],
) -> bool:
    ok = True
    exception_map = build_exception_map(special_schedules)

    print(f"\n[검증] {result.year}년 {result.month}월 근무표")

    # 1) 상담사별 1번 좌석 배정 횟수 (형평성 확인, 제외 대상 제외)
    eligible = [c for c in all_counselors if c not in seat1_excluded]
    counts = {c: result.seat1_counts_this_month.get(c, 0) for c in eligible}
    print("  1) 상담사별 1번 좌석 배정 횟수(이번 달):")
    for c, cnt in counts.items():
        print(f"     - {c}: {cnt}회")
    if counts:
        diff = max(counts.values()) - min(counts.values())
        status = "OK" if diff <= 1 else "FAIL"
        print(f"     -> 최대-최소 차이: {diff} [{status}]")
        ok = ok and diff <= 1

    # 2) 제외 대상이 1번 좌석에 배정되지 않았는지
    violations = [
        (d, s.seat1) for d, s in result.days.items() if s.seat1 in seat1_excluded
    ]
    status = "OK" if not violations else "FAIL"
    print(f"  2) 제외 대상 1번 좌석 배정 여부: [{status}]")
    if violations:
        ok = False
        for d, name in violations:
            print(f"     - 위반: {d} {name}")

    # 3) 근무일마다 정확히 2명 배정되었는지
    missing = []
    for d, s in result.days.items():
        if s.note is not None:
            continue
        if not s.seat1 or not s.seat2 or s.seat1 == s.seat2:
            missing.append(d)
    status = "OK" if not missing else "FAIL"
    print(f"  3) 근무일별 2명 배정 여부: [{status}]")
    if missing:
        ok = False
        for d in missing:
            print(f"     - 위반: {d}")

    # 4) 예외일에 근무자가 배정되지 않았는지
    exception_violations = [
        d for d in exception_map
        if d in result.days and (result.days[d].seat1 or result.days[d].seat2)
    ]
    status = "OK" if not exception_violations else "FAIL"
    print(f"  4) 예외일 근무 미배정 여부: [{status}]")
    if exception_violations:
        ok = False
        for d in exception_violations:
            print(f"     - 위반: {d}")

    print(f"  => 종합 결과: {'모두 통과' if ok else '문제 발견'}")
    return ok


# =============================================================================
# 엑셀(.xlsx) 출력
# =============================================================================

SEAT1_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def export_to_excel(
    result: ScheduleResult,
    all_counselors: list[str],
    seat1_excluded: list[str],
    output_path: str,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = f"{result.month}월 근무표"

    n_cols = len(WEEKDAY_LABELS)
    for col in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16

    # 제목
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    title_cell = ws.cell(row=1, column=1, value=f"[{result.month}월 상담센터 근무 순서]")
    title_cell.font = Font(size=14, bold=True)
    title_cell.alignment = CENTER

    row = 3
    # 요일 헤더
    for col, label in enumerate(WEEKDAY_LABELS, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = Font(bold=True)
        cell.alignment = CENTER
        cell.border = THIN_BORDER
    row += 1

    weeks = calendar.monthcalendar(result.year, result.month)
    for week in weeks:
        weekday_dates = [
            date(result.year, result.month, week[i]) if week[i] != 0 else None
            for i in range(5)
        ]
        if all(d is None for d in weekday_dates):
            continue

        # 날짜 행
        for col, d in enumerate(weekday_dates, start=1):
            cell = ws.cell(row=row, column=col, value=f"{result.month}월 {d.day}일" if d else "")
            cell.font = Font(bold=True)
            cell.alignment = CENTER
            cell.border = THIN_BORDER
        date_row = row
        row += 1

        # 근무자 행 (예외일은 두 행을 병합해서 사유 표시)
        seat1_row = row
        seat2_row = row + 1
        for col, d in enumerate(weekday_dates, start=1):
            c1 = ws.cell(row=seat1_row, column=col)
            c2 = ws.cell(row=seat2_row, column=col)
            c1.border = THIN_BORDER
            c2.border = THIN_BORDER
            c1.alignment = CENTER
            c2.alignment = CENTER

            if d is None:
                continue

            day_info = result.days.get(d)
            if day_info is None:
                continue

            if day_info.note is not None:
                ws.merge_cells(start_row=seat1_row, start_column=col, end_row=seat2_row, end_column=col)
                c1.value = day_info.note
                c1.alignment = CENTER
            else:
                c1.value = day_info.seat1
                c1.fill = SEAT1_FILL
                c2.value = day_info.seat2
        row = seat2_row + 1

    row += 1
    footnote = f"※ 시니어일자리 지원센터 컨설턴트 : {', '.join(seat1_excluded)}"
    ws.cell(row=row, column=1, value=footnote).font = Font(italic=True)

    wb.save(output_path)
    print(f"\n엑셀 파일 저장 완료: {output_path}")


# =============================================================================
# 실행부
# =============================================================================

def main() -> None:
    saved_counts, saved_last_date = load_state(STATE_FILE)
    prev_counts = {**saved_counts, **PREV_SEAT1_COUNTS}
    prev_last_date = {**saved_last_date, **PREV_SEAT1_LAST_DATE}

    result = generate_schedule(
        year=YEAR,
        month=MONTH,
        all_counselors=ALL_COUNSELORS,
        seat1_excluded=SEAT1_EXCLUDED,
        special_schedules=SPECIAL_SCHEDULES,
        prev_seat1_counts=prev_counts,
        prev_seat1_last_date=prev_last_date,
    )

    print(f"[{result.month}월 상담센터 근무 순서]\n")
    for d in sorted(result.days):
        s = result.days[d]
        weekday_label = WEEKDAY_LABELS[d.weekday()]
        if s.note is not None:
            print(f"{d} ({weekday_label}): {s.note}")
        else:
            marker = " *" if s.seat1 else ""
            print(f"{d} ({weekday_label}): 1번={s.seat1}{marker} / 2번={s.seat2}")

    ok = validate_schedule(result, ALL_COUNSELORS, SEAT1_EXCLUDED, SPECIAL_SCHEDULES)
    export_to_excel(result, ALL_COUNSELORS, SEAT1_EXCLUDED, OUTPUT_FILENAME)

    if ok:
        save_state(STATE_FILE, result.seat1_counts_cumulative, result.seat1_last_date)
        print(f"다음 달 형평성 이어받기용 상태 저장 완료: {STATE_FILE}")
    else:
        print("검증에 실패해 상태 파일은 갱신하지 않았습니다. 설정값을 확인해 주세요.")


if __name__ == "__main__":
    main()
